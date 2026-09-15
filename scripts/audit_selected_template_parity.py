#!/usr/bin/env python3
"""Offline only: inventory ETALON differences and score source-annotated facts.

Nothing here is imported by the filling application. Completed example values
are comparison data, never assignments or model context. An unreviewed cell is
not classified as missing from PDF, and a partial score is not overall parity.
"""
from __future__ import annotations

import argparse
import json
import sys
from collections import Counter
from datetime import date, datetime
from pathlib import Path

import openpyxl
import yaml

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
from executive_docs.excel import sha256
from executive_docs.evidence_matching import normalize_evidence_text
from executive_docs.selected_templates import TemplateCatalog


def value_json(value):
    return value.isoformat() if isinstance(value, (date, datetime)) else value


def audit(catalog, template_id: str, pdf: Path, *, output: Path | None = None, case_path: Path | None = None) -> dict:
    contract = catalog.get(template_id)
    case = yaml.safe_load(case_path.read_text()) if case_path else {}
    if case:
        if case["template_id"] != template_id or case["pdf_sha256"] != sha256(pdf):
            raise ValueError("Контрольный случай относится к другому шаблону или PDF")
        if case["etalon_sha256"] != contract.etalon_sha256:
            raise ValueError("Контрольный ETALON изменился; нужна повторная проверка")
    annotations = {tuple(item["target"]): item for item in case.get("checks", [])}
    metadata = yaml.safe_load(contract.path.read_text())
    source_path = catalog.root / metadata["source_template"]
    etalon_path = catalog.root / metadata["etalon_example"]
    if sha256(etalon_path) != contract.etalon_sha256 or sha256(source_path) != contract.source_sha256:
        raise ValueError("Исходный шаблон или ETALON изменился")
    books = [openpyxl.load_workbook(path, data_only=False) for path in (source_path, etalon_path, catalog.candidate_path(contract))]
    source, etalon, candidate = books
    result = openpyxl.load_workbook(output, data_only=False) if output else None
    if result:
        books.append(result)
    try:
        inventory = []
        targets = set(contract.field_map) | set(annotations)
        # Include changed ETALON literals absent from the current input map.
        # Do not treat unchanged sample values as evidence for this project.
        unmapped_changes = []
        for sheet in etalon:
            if sheet.title not in source.sheetnames:
                continue
            for row in sheet:
                for cell in row:
                    if cell.value not in (None, "") and cell.data_type != "f" and cell.value != source[sheet.title][cell.coordinate].value:
                        target = sheet.title, cell.coordinate
                        targets.add(target)
                        if target not in contract.field_map:
                            unmapped_changes.append(target)
        for target in sorted(targets):
            sheet, cell = target
            field = contract.field_map.get(target)
            expected = etalon[sheet][cell].value if sheet in etalon.sheetnames else None
            original = source[sheet][cell].value if sheet in source.sheetnames else None
            annotation = annotations.get(target)
            if annotation:
                classification = annotation["classification"]
            elif isinstance(expected, str) and expected.startswith("="):
                classification = "formula_not_independent_fact"
            elif expected not in (None, "") and expected == original:
                classification = "unchanged_example_needs_source_review"
            elif field and field.evidence_rule in {"actual_executive_document_only", "signatory_role_pdf", "authority_document_pdf"}:
                classification = "requires_execution_or_role_evidence"
            else:
                classification = "unreviewed"
            formula = candidate[sheet][cell].value if sheet in candidate.sheetnames else None
            actual = result[sheet][cell].value if result and sheet in result.sheetnames else None
            check = None
            if annotation and result:
                if classification == "pdf_supported":
                    # Explicit variants only; do not silently accept extra facts.
                    check = normalize_evidence_text(str(actual or "")) in {
                        normalize_evidence_text(value) for value in annotation["accepted_values"]
                    }
                elif classification == "requires_additional_evidence":
                    check = actual in (None, "")
            inventory.append({
                "sheet": sheet, "cell": cell, "semantic_id": field.semantic_id if field else None,
                "classification": classification, "model_writable": bool(field and not field.manual_reason),
                "etalon_value": value_json(expected), "output_value": value_json(actual),
                "candidate_formula": formula if isinstance(formula, str) and formula.startswith("=") else None,
                "annotation": annotation, "check_passed": check,
            })
        scored = [item for item in inventory if item["classification"] == "pdf_supported" and item["annotation"]]
        checked = [item for item in inventory if item["check_passed"] is not None]
        return {
            "template_id": template_id, "template_version": contract.version,
            "pdf_sha256": sha256(pdf), "output_sha256": sha256(output) if output else None,
            "scope": "offline_partial_source_annotated_benchmark_not_overall_etalon_parity",
            "overall_etalon_coverage": None,
            "reviewed_pdf_targets": len(scored),
            "matched_reviewed_pdf_targets": sum(item["check_passed"] is True for item in scored) if result else None,
            "failed_checks": sum(item["check_passed"] is False for item in checked) if result else None,
            "classification_counts": dict(Counter(item["classification"] for item in inventory)),
            "unmapped_etalon_changes": unmapped_changes,
            "inventory": inventory,
        }
    finally:
        for book in books:
            book.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("pdf", type=Path)
    parser.add_argument("--template-id", required=True)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--case", type=Path)
    parser.add_argument("--report-dir", type=Path, required=True)
    args = parser.parse_args()
    destination = args.report_dir.resolve()
    if not destination.is_relative_to(ROOT / "data/runs"):
        raise SystemExit("Отчёт можно записать только под data/runs/")
    catalog = TemplateCatalog(ROOT, ROOT / "templates/fill-contracts", ROOT / "templates/approved")
    report = audit(catalog, args.template_id, args.pdf, output=args.output, case_path=args.case)
    destination.mkdir(parents=True, exist_ok=True)
    (destination / "semantic-parity.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps({key: value for key, value in report.items() if key != "inventory"}, ensure_ascii=False, indent=2))
    return 1 if report["failed_checks"] else 0


if __name__ == "__main__":
    raise SystemExit(main())
