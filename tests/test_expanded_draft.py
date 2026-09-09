from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest
import yaml

from executive_docs.agent import OpenAIAgent
from executive_docs.domain import Artifact, ProjectState, TemplateCellAssignment
from executive_docs.excel import MAIN_NS, OOXMLWorkbook, sha256
from executive_docs.selected_templates import (
    PROJECT_DRAFT_HEADER,
    PROJECT_FILL_RGB,
    SelectedTemplateGenerator,
    TemplateCatalog,
    validate_selected_template_output,
)


SHEET = "Данные"
PROJECT_QUOTE = "Проектное количество провода: 12,5 м."
ACTUAL_QUOTE = "Акт выполненных работ: фактически смонтировано 12,5 м провода."
DATE_QUOTE = "Плановая дата начала работ: 01.07.2026."
SIGNER_QUOTE = "Проектировщик ООО «Пример», директор Иванов И.И."


def draft_catalog(tmp_path: Path, *, materials: bool = False) -> TemplateCatalog:
    approved = tmp_path / "templates/approved"
    contracts = tmp_path / "templates/fill-contracts"
    approved.mkdir(parents=True)
    contracts.mkdir(parents=True)
    candidate = approved / "expanded.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = SHEET
    sheet["A1"] = "Количество провода"
    sheet["C1"] = "=B1"
    sheet["B1"].number_format = "0.00"
    sheet.oddHeader.right.text = "Исходный колонтитул"
    support = workbook.create_sheet("Служебный")
    support["A1"] = "='Данные'!B1"
    support.sheet_state = "hidden"
    workbook.save(candidate)
    workbook.close()

    def field(cell: str, semantic_id: str, **overrides: object) -> dict:
        return {
            "sheet": SHEET,
            "cell": cell,
            "label": semantic_id,
            "semantic_id": semantic_id,
            "value_kind": "text",
            "evidence_rule": "direct_pdf",
            "required": True,
            **overrides,
        }

    contract = {
        "template_id": "expanded",
        "display_name": "Тест расширенного черновика",
        "document_kind": "test",
        "version": "1-candidate",
        "status": "DISCOVERY_REVIEW_REQUIRED",
        "approved": False,
        "candidate_template": "templates/approved/expanded.xlsx",
        "candidate_sha256": sha256(candidate),
        "source_sha256": "0" * 64,
        "etalon_sha256": "1" * 64,
        "output_filename": "expanded.xlsx",
        "warning_fill_rgb": "FFFFE699",
        "structural_findings": {},
        "fields": [
            field("B1", "aosr.wire.quantity", value_kind="number",
                  evidence_rule="actual_executive_document_only", allow_project_basis=True),
            field("B2", "aosr.other.quantity", value_kind="number",
                  evidence_rule="actual_executive_document_only"),
            field("B3", "actual.start_date", evidence_rule="actual_executive_document_only"),
            field("B4", "contractor.site_representative.name", evidence_rule="signatory_role_pdf"),
            field("B5", "project.code"),
            field("B6", "project.explicitly_missing"),
            field("B7", "project.omitted"),
            field("B8", "project.rejected"),
            field("B9", "unknown.mapping", manual_reason="Назначение требует проверки"),
            field("B10", "aosr.wire.repeated_quantity", value_kind="number",
                  evidence_rule="actual_executive_document_only", allow_project_basis=True),
            field("B11", "designer.name", evidence_rule="organization_role_pdf"),
            field("B12", "designer.address", evidence_rule="organization_role_pdf"),
        ],
    }
    if materials:
        contract["fields"].extend([
            field("D1", "test.materials.item_1.name", allow_project_basis=True),
            field("E1", "test.materials.item_1.type", allow_project_basis=True),
            field("F1", "test.materials.item_1.quantity", value_kind="number", allow_project_basis=True),
        ])
    (contracts / "expanded.yaml").write_text(
        yaml.safe_dump(contract, allow_unicode=True, sort_keys=False), encoding="utf-8",
    )
    return TemplateCatalog(tmp_path, contracts, approved)


def assignment(cell: str = "B1", *, value: str = "12,5", basis: str = "project",
               quote: str = PROJECT_QUOTE) -> dict:
    return {
        "sheet": SHEET,
        "cell": cell,
        "value": value,
        "source_file_id": "project",
        "locator": "page:1",
        "evidence_fragment": quote,
        "value_basis": basis,
    }


@pytest.fixture
def indexed_context(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    """Exercise validation against a controlled reliable index, not visual trust."""
    catalog = draft_catalog(tmp_path)
    artifact = Artifact(
        id="project", original_name="project.pdf", stored_name="project.pdf",
        media_type="application/pdf", size=1, sha256="2" * 64, pages=1,
    )
    state = ProjectState(job_id="expanded-test", operator_name="test", artifacts=[artifact])
    text = "\n".join((PROJECT_QUOTE, ACTUAL_QUOTE, DATE_QUOTE, SIGNER_QUOTE, "Шифр P-42"))
    monkeypatch.setattr("executive_docs.agent.source_index", lambda *args, **kwargs: {
        "segments": [{"page": 1, "text": text, "text_reliable": True, "visual_required": False}],
    })
    return catalog, catalog.get("expanded"), state, tmp_path


@pytest.mark.parametrize("cell,basis,quote,expected", [
    ("B1", "project", PROJECT_QUOTE, True),
    ("B1", "document", PROJECT_QUOTE, False),
    ("B2", "project", PROJECT_QUOTE, False),
    ("B2", "document", ACTUAL_QUOTE, True),
    ("B1", "document", ACTUAL_QUOTE, True),
])
def test_quantity_requires_project_permission_and_honest_basis(
    indexed_context, cell: str, basis: str, quote: str, expected: bool,
) -> None:
    _, contract, state, root = indexed_context
    result, diagnostics = OpenAIAgent._recover_template_fill(
        state, contract, [{"assignments": [assignment(cell, basis=basis, quote=quote)]}], root,
    )
    assert bool(result.assignments) is expected
    assert bool(diagnostics) is not expected
    if expected:
        assert result.assignments[0].value_basis == basis
    else:
        finding = next(item for item in result.unresolved if item.cell == cell)
        assert finding.category == "rejected"


@pytest.mark.parametrize("cell,value,quote", [
    ("B3", "01.07.2026", DATE_QUOTE),
    ("B4", "Иванов И.И.", SIGNER_QUOTE),
])
def test_project_basis_cannot_fill_actual_dates_or_execution_signers(
    indexed_context, cell: str, value: str, quote: str,
) -> None:
    _, contract, state, root = indexed_context
    assert not contract.field_map[(SHEET, cell)].allow_project_basis
    result, diagnostics = OpenAIAgent._recover_template_fill(
        state, contract,
        [{"assignments": [assignment(cell, value=value, quote=quote)]}], root,
    )
    assert not result.assignments
    assert diagnostics
    assert next(item for item in result.unresolved if item.cell == cell).category == "rejected"


def test_project_basis_still_requires_value_on_cited_page(indexed_context) -> None:
    _, contract, state, root = indexed_context
    result, diagnostics = OpenAIAgent._recover_template_fill(
        state, contract, [{"assignments": [assignment(value="125")]}], root,
    )
    assert not result.assignments
    assert "value is not present" in diagnostics[0]["reason"]


def test_one_source_fact_can_fill_multiple_matching_registered_targets(indexed_context) -> None:
    _, contract, state, root = indexed_context
    result, diagnostics = OpenAIAgent._recover_template_fill(
        state, contract, [{"assignments": [assignment(), assignment("B10")]}], root,
    )
    assert {item.cell for item in result.assignments} == {"B1", "B10"}
    assert {item.locator for item in result.assignments} == {"page:1"}
    assert not diagnostics


def test_recovery_distinguishes_omission_rejection_and_explicit_absence(indexed_context) -> None:
    _, contract, state, root = indexed_context
    payload = {
        "assignments": [
            assignment("B5", value="P-42", basis="document", quote="Шифр P-42"),
            assignment("B8", value="P-77", basis="document", quote="Шифр P-42"),
        ],
        "unresolved": [{
            "sheet": SHEET, "cell": "B6", "category": "missing_from_pdf",
            "reason": "В документе нет такого реквизита",
        }],
    }
    result, diagnostics = OpenAIAgent._recover_template_fill(state, contract, [payload], root)
    findings = {item.cell: item for item in result.unresolved}
    assert [item.cell for item in result.assignments] == ["B5"]
    assert findings["B6"].category == "missing_from_pdf"
    assert findings["B7"].category == "not_returned"
    assert findings["B8"].category == "rejected"
    assert "value is not present" in diagnostics[0]["reason"]
    unresolved = SelectedTemplateGenerator.unresolved_cells(contract, result.assignments, result.unresolved)
    assert next(item for item in unresolved if item.cell == "B9").category == "manual_confirmation"


def generated_project_draft(tmp_path: Path):
    catalog = draft_catalog(tmp_path)
    contract = catalog.get("expanded")
    assignments = [TemplateCellAssignment.model_validate(assignment())]
    source = catalog.candidate_path(contract)
    original = source.read_bytes()
    output, unresolved = SelectedTemplateGenerator(catalog).generate(contract, assignments, tmp_path / "output")
    assert source.read_bytes() == original
    return catalog, contract, assignments, source, output, unresolved


def test_project_draft_keeps_numeric_values_distinct_markers_and_all_print_headers(tmp_path: Path) -> None:
    _, contract, assignments, source, output, unresolved = generated_project_draft(tmp_path)
    book = openpyxl.load_workbook(output, data_only=False)
    try:
        assert book[SHEET]["B1"].value == 12.5
        assert book[SHEET]["B1"].data_type == "n"
        assert book[SHEET]["B1"].number_format == "0.00"
        assert book[SHEET]["B1"].fill.fgColor.rgb == PROJECT_FILL_RGB
        assert book[SHEET]["B2"].value is None
        assert book[SHEET]["B2"].fill.fgColor.rgb == contract.warning_fill_rgb
        assert book[SHEET]["C1"].value == "=B1"
        assert book["Служебный"].sheet_state == "hidden"
        assert book[SHEET].oddHeader.right.text == "Исходный колонтитул"
    finally:
        book.close()
    archive = OOXMLWorkbook(output)
    for name in archive.sheet_parts:
        root = archive._sheet_root(name)
        for header in ("oddHeader", "evenHeader", "firstHeader"):
            assert root.find(f"{{{MAIN_NS}}}headerFooter/{{{MAIN_NS}}}{header}").text.startswith(PROJECT_DRAFT_HEADER)
    assert (SHEET, "B1") not in {(item.sheet, item.cell) for item in unresolved}
    issues = validate_selected_template_output(output, source, contract, assignments, unresolved)
    assert not [item for item in issues if item.severity == "error"]


@pytest.mark.parametrize("tamper", ["project_fill", "unresolved_fill", "project_header"])
def test_tampering_with_project_or_unresolved_markers_is_detected(tmp_path: Path, tamper: str) -> None:
    _, contract, assignments, source, output, unresolved = generated_project_draft(tmp_path)
    archive = OOXMLWorkbook(output)
    if tamper == "project_fill":
        archive.highlight_cells({SHEET: ["B1"]}, rgb=contract.warning_fill_rgb)
    elif tamper == "unresolved_fill":
        archive.highlight_cells({SHEET: ["B2"]}, rgb=PROJECT_FILL_RGB)
    else:
        root = archive._sheet_root("Служебный")
        root.find(f"{{{MAIN_NS}}}headerFooter/{{{MAIN_NS}}}evenHeader").text = ""
        archive._save_sheet_root("Служебный", root)
    archive.save(output)
    issues = validate_selected_template_output(output, source, contract, assignments, unresolved)
    errors = {item.code for item in issues if item.severity == "error"}
    assert errors
    if tamper == "project_fill":
        assert "PROJECT_CELL_NOT_HIGHLIGHTED" in errors
    elif tamper == "unresolved_fill":
        assert "UNRESOLVED_CELL_NOT_HIGHLIGHTED" in errors
    else:
        assert "WORKSHEET_STRUCTURE_CHANGED" in errors


def test_document_only_draft_does_not_add_project_marker(tmp_path: Path) -> None:
    catalog = draft_catalog(tmp_path)
    contract = catalog.get("expanded")
    assignments = [TemplateCellAssignment.model_validate(assignment(basis="document", quote=ACTUAL_QUOTE))]
    output, unresolved = SelectedTemplateGenerator(catalog).generate(contract, assignments, tmp_path / "output")
    archive = OOXMLWorkbook(output)
    for name in archive.sheet_parts:
        assert PROJECT_DRAFT_HEADER not in archive._sheet_root(name).xpath("string()")
    issues = validate_selected_template_output(output, catalog.candidate_path(contract), contract, assignments, unresolved)
    assert not [item for item in issues if item.severity == "error"]


def test_a_wire_dimension_cannot_be_used_as_the_material_quantity(indexed_context, monkeypatch) -> None:
    _, contract, state, root = indexed_context
    quote = "СИП-2 3×70+1×95 — 308 м"
    monkeypatch.setattr("executive_docs.agent.source_index", lambda *args, **kwargs: {
        "segments": [{"page": 1, "text": quote, "text_reliable": True, "visual_required": False}],
    })
    result, diagnostics = OpenAIAgent._recover_template_fill(
        state, contract, [{"assignments": [assignment(value="70", quote=quote)]}], root,
    )
    assert not result.assignments
    assert diagnostics
    assert next(item for item in result.unresolved if item.cell == "B1").category == "rejected"


def test_a_clipped_quote_cannot_disguise_part_of_a_decimal_quantity(indexed_context, monkeypatch) -> None:
    _, contract, state, root = indexed_context
    monkeypatch.setattr("executive_docs.agent.source_index", lambda *args, **kwargs: {
        "segments": [{"page": 1, "text": "Длина12,5 м", "text_reliable": True, "visual_required": False}],
    })
    result, diagnostics = OpenAIAgent._recover_template_fill(
        state, contract, [{"assignments": [assignment(value="5", quote="5 м")]}], root,
    )
    assert not result.assignments
    assert diagnostics
    assert next(item for item in result.unresolved if item.cell == "B1").category == "rejected"


@pytest.mark.parametrize("other_name", ["АБ", "А Строй"])
@pytest.mark.parametrize("target", ["name", "address"])
def test_normalization_does_not_borrow_another_organizations_name_or_role(
    indexed_context, monkeypatch, other_name: str, target: str,
) -> None:
    _, contract, state, root = indexed_context
    role_quote = f"Проектная организация: ООО «{other_name}»"
    if target == "name":
        proposed = assignment("B11", value="ООО «А»", quote=role_quote, basis="document")
        source_text = role_quote
    else:
        address_quote = "ООО «А». Адрес: Москва"
        proposed = assignment("B12", value="Москва", quote=address_quote, basis="document")
        proposed["subject_name"] = "ООО «А»"
        proposed["context_evidence"] = [{"locator": "page:1", "evidence_fragment": role_quote}]
        source_text = address_quote + "\n" + role_quote
    monkeypatch.setattr("executive_docs.agent.source_index", lambda *args, **kwargs: {
        "segments": [{"page": 1, "text": source_text, "text_reliable": True, "visual_required": False}],
    })
    result, diagnostics = OpenAIAgent._recover_template_fill(
        state, contract, [{"assignments": [proposed]}], root,
    )
    assert not result.assignments
    assert diagnostics
    assert next(item for item in result.unresolved if item.cell == proposed["cell"]).category == "rejected"


MATERIAL_QUOTE = "Провод, СИП-2 3×70+1×95 — 308 м"


def material_assignments() -> list[TemplateCellAssignment]:
    return [TemplateCellAssignment.model_validate(assignment(cell, value=value, quote=MATERIAL_QUOTE))
            for cell, value in (("D1", "Провод"), ("E1", "СИП-2 3×70+1×95"), ("F1", "308"))]


@pytest.mark.parametrize("part", ["type", "quantity"])
def test_material_type_or_quantity_without_its_name_cannot_generate(tmp_path: Path, part: str) -> None:
    catalog = draft_catalog(tmp_path, materials=True)
    contract = catalog.get("expanded")
    proposed = material_assignments()[1 if part == "type" else 2]
    with pytest.raises(ValueError, match="одной позиции PDF"):
        SelectedTemplateGenerator(catalog).generate(contract, [proposed], tmp_path / "output")


@pytest.mark.parametrize("changed_evidence", [
    {"source_file_id": "another-pdf"},
    {"locator": "page:2"},
    {"evidence_fragment": "Другая позиция: арматура — 308 м"},
])
def test_material_columns_must_share_source_page_and_quote(tmp_path: Path, changed_evidence: dict) -> None:
    catalog = draft_catalog(tmp_path, materials=True)
    contract = catalog.get("expanded")
    proposed = material_assignments()
    proposed[-1] = proposed[-1].model_copy(update=changed_evidence)
    with pytest.raises(ValueError, match="одной позиции PDF"):
        SelectedTemplateGenerator(catalog).generate(contract, proposed, tmp_path / "output")


@pytest.mark.parametrize("include_type", [False, True])
def test_material_name_and_quantity_with_shared_normalized_evidence_generate(
    tmp_path: Path, include_type: bool,
) -> None:
    catalog = draft_catalog(tmp_path, materials=True)
    contract = catalog.get("expanded")
    proposed = material_assignments()
    proposed[-1] = proposed[-1].model_copy(update={"evidence_fragment": MATERIAL_QUOTE.replace(",", ";")})
    if not include_type:
        proposed = [proposed[0], proposed[-1]]
    output, unresolved = SelectedTemplateGenerator(catalog).generate(contract, proposed, tmp_path / "output")
    book = openpyxl.load_workbook(output, data_only=False)
    try:
        assert book[SHEET]["D1"].value == "Провод"
        assert book[SHEET]["F1"].value == 308
        assert book[SHEET]["F1"].data_type == "n"
    finally:
        book.close()
    issues = validate_selected_template_output(output, catalog.candidate_path(contract), contract, proposed, unresolved)
    assert not [item for item in issues if item.severity == "error"]
