from __future__ import annotations

import importlib.util
from pathlib import Path

import openpyxl
import pytest

from executive_docs.selected_templates import TemplateCatalog


ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location(
    "register_selected_templates", ROOT / "scripts" / "register_selected_templates.py"
)
assert spec is not None and spec.loader is not None
registration = importlib.util.module_from_spec(spec)
spec.loader.exec_module(registration)


def test_ojr_main_installation_and_printed_roles_are_not_misrouted():
    catalog = TemplateCatalog(ROOT, ROOT / "templates/fill-contracts", ROOT / "templates/approved")
    contract = catalog.get("ojr")
    assert contract.field_map[("Данные объект", "B11")].semantic_id == "project.installation.primary"
    assert contract.field_map[("Титульный лист", "A13")].semantic_id == "developer.name"
    assert contract.field_map[("Титульный лист", "Q4")].semantic_id == "actual.journal_number"
    book = openpyxl.load_workbook(catalog.candidate_path(contract))
    try:
        title = book["Титульный лист"]
        assert "B11" in title["A178"].value
        assert "B17" in title["A85"].value
        assert "B18" in title["A87"].value and "B19" in title["A87"].value
        assert "B32" in title["A31"].value
        assert "B52" in title["A67"].value
        for coordinate in ("Q4", "A13", "A15", "A17", "A33", "A69", "A89"):
            assert title[coordinate].value is None
            assert ("Титульный лист", coordinate) in contract.field_map
        assert book["Данные объект"]["C33"].value.startswith('=IF(B33="","",')
        assert not book["Данные объект"].row_dimensions[13].hidden
    finally:
        book.close()


def test_ojr_repair_refuses_changed_source_formula():
    source = openpyxl.load_workbook(registration.SOURCE_DIR / "3. ОЖР.xlsx")
    source["Титульный лист"]["A69"] = "='Данные объект'!B51"
    with pytest.raises(ValueError, match="Изменилась проверяемая связь"):
        registration.ojr_layout_changes(source)
    source.close()


@pytest.mark.parametrize("template_id,filename,count", [("emr", "1. ЭМР1.xlsx", 20), ("avk", "4. АВК.xlsx", 2)])
def test_irrecoverable_formula_quarantine_is_explicit_blocking_and_blank(
    tmp_path: Path, template_id: str, filename: str, count: int,
) -> None:
    from executive_docs.selected_templates import SelectedTemplateGenerator, validate_selected_template_output

    catalog = TemplateCatalog(ROOT, ROOT / "templates/fill-contracts", ROOT / "templates/approved")
    contract = catalog.get(template_id)
    source = openpyxl.load_workbook(registration.SOURCE_DIR / filename)
    candidate = openpyxl.load_workbook(catalog.candidate_path(contract))
    try:
        quarantined = registration.broken_formula_quarantine(template_id, source)
        assert len(quarantined) == count
        assert contract.structural_findings["quarantined_broken_formula_count"] == count
        for key in ("formula_errors_observed", "raw_ref_error_count", "remaining_external_formula_reference_count", "unsafe_blank_formula_count"):
            assert contract.structural_findings[key] == 0
        for (sheet, cell) in quarantined:
            field = contract.field_map[(sheet, cell)]
            assert field.required and field.manual_reason and not field.allow_project_basis
            assert "утрачена ссылка" in field.manual_reason
            assert field.evidence_rule == "actual_executive_document_only"
            assert candidate[sheet][cell].value is None
            assert candidate[sheet][cell]._style == source[sheet][cell]._style
        # The two broken EMR roots have these eighteen healthy downstream
        # links. Leave them blank-safe and live; never replace them with values.
        if template_id == "emr":
            for row in range(68, 85):
                assert candidate["ВЛ-1"][f"F{row}"].value == '=IF(BG25="","",BG25)'
            assert candidate["24 ВЛИ"]["G50"].value == '=IF(G29="","",G29)'
    finally:
        candidate.close()
        source.close()
    output, unresolved = SelectedTemplateGenerator(catalog).generate(contract, [], tmp_path)
    unresolved_map = {(item.sheet, item.cell): item for item in unresolved}
    book = openpyxl.load_workbook(output)
    try:
        for sheet, cell in quarantined:
            item = unresolved_map[(sheet, cell)]
            assert item.category == "manual_confirmation" and item.blocking
            assert book[sheet][cell].value is None
            assert book[sheet][cell].fill.fgColor.rgb == contract.warning_fill_rgb
    finally:
        book.close()
    issues = validate_selected_template_output(output, catalog.candidate_path(contract), contract, [], unresolved)
    assert not [issue for issue in issues if issue.severity == "error"]


@pytest.mark.parametrize("template_id,filename", [("emr", "1. ЭМР1.xlsx"), ("avk", "4. АВК.xlsx")])
def test_formula_quarantine_does_not_silently_clear_a_repaired_source(
    template_id: str, filename: str,
) -> None:
    source = openpyxl.load_workbook(registration.SOURCE_DIR / filename)
    try:
        targets = registration.broken_formula_quarantine(template_id, source)
        sheet, cell = next(iter(targets))
        source[sheet][cell] = "='Данные объект'!B3"
        with pytest.raises(ValueError, match="Изменилась карантинируемая формула"):
            registration.broken_formula_quarantine(template_id, source)
    finally:
        source.close()


@pytest.mark.parametrize("template_id", ["emr", "protocols", "ojr", "avk", "aosr_vl"])
def test_all_pdf_writable_rows_are_visible_and_blank_guards_applied(template_id):
    catalog = TemplateCatalog(ROOT, ROOT / "templates/fill-contracts", ROOT / "templates/approved")
    contract = catalog.get(template_id)
    assert contract.structural_findings["blank_preserving_formula_count"] > 0
    book = openpyxl.load_workbook(catalog.candidate_path(contract))
    try:
        for field in contract.fields:
            if not field.manual_reason:
                assert not book[field.sheet].row_dimensions[book[field.sheet][field.cell].row].hidden, field.coordinate
        headers = {"emr": ("Ведомость общ", ["E17", "AI17", "BE17"]), "avk": (" Журнал АВК", ["C34", "D34"])}
        if template_id in headers:
            sheet, coordinates = headers[template_id]
            for coordinate in coordinates:
                assert book[sheet][coordinate].value
                assert (sheet, coordinate) not in contract.field_map
    finally:
        book.close()


@pytest.mark.parametrize("template_id", ["emr", "protocols", "ojr", "avk", "aosr_vl"])
def test_object_card_and_actual_dates_are_eligible_with_correct_pdf_evidence(
    template_id: str,
) -> None:
    catalog = TemplateCatalog(
        ROOT, ROOT / "templates" / "fill-contracts", ROOT / "templates" / "approved"
    )
    contract = catalog.get(template_id)
    fields = contract.field_map
    for cell, semantic_id in {
        "B2": "project.sap_number",
        "B3": "project.object_name",
        "B4": "project.district",
        "B5": "project.object_address",
    }.items():
        field = fields[("Данные объект", cell)]
        assert field.semantic_id == semantic_id
        assert field.manual_reason is None
        assert field.evidence_rule == "direct_pdf"
    sap = fields[("Данные объект", "B2")]
    sap.validate_raw_value("SAP - 123456")
    with pytest.raises(ValueError, match="смыслу поля"):
        sap.validate_raw_value("5557-354783-68-03/26")

    actual_dates = ("B6", "B7") if template_id in {"avk", "aosr_vl"} else ("B7", "B8")
    for cell, semantic_id in zip(actual_dates, ("actual.start", "actual.end")):
        field = fields[("Данные объект", cell)]
        assert field.manual_reason is None
        assert field.semantic_id == semantic_id
        assert field.value_kind == "date"
        assert field.evidence_rule == "actual_executive_document_only"
        assert "проектный график" in field.description
    assert contract.status == "DISCOVERY_REVIEW_REQUIRED"
    assert contract.approved is False


@pytest.mark.parametrize("template_id,filename,_,__", registration.TEMPLATES)
def test_object_card_mapping_comes_from_source_captions_and_rejects_layout_drift(
    template_id: str, filename: str, _: str, __: str,
) -> None:
    source = openpyxl.load_workbook(registration.SOURCE_DIR / filename)
    try:
        overrides = registration.object_card_field_overrides(template_id, source)
        assert ("Данные объект", "B5") in overrides
        # A changed caption may mean the coordinate now holds an organization
        # address. Refuse the mapping rather than silently expanding model access.
        source["Данные объект"]["A5"] = "Адрес подрядчика"
        with pytest.raises(ValueError, match="Изменилась подпись"):
            registration.object_card_field_overrides(template_id, source)
    finally:
        source.close()


@pytest.mark.parametrize("template_id,organization_row,customer_row,designer_row", [
    ("aosr_vl", 11, 23, 39), ("avk", 11, 23, 36),
    ("ojr", 17, 30, 50), ("protocols", 22, None, 50),
    ("emr", 19, 33, None),
])
def test_organization_facts_are_role_bound_and_do_not_require_profiles(
    template_id: str, organization_row: int, customer_row: int | None,
    designer_row: int | None,
) -> None:
    catalog = TemplateCatalog(ROOT, ROOT / "templates/fill-contracts", ROOT / "templates/approved")
    contract = catalog.get(template_id)
    for role, row in (("contractor", organization_row), ("customer", customer_row), ("designer", designer_row)):
        if row is None:
            continue
        field = contract.field_map[("Данные объект", f"B{row}")]
        assert field.semantic_id == f"{role}.name"
        assert field.manual_reason is None
        assert field.evidence_rule == "organization_role_pdf"
        assert "противоречащих организациях" in field.description
        assert "при необходимости сослаться" in field.description
    for field in contract.fields:
        assert "профил" not in (field.manual_reason or "").casefold()
        assert "отсутствует в pdf" not in (field.manual_reason or "").casefold()
        if field.manual_reason is None:
            assert field.semantic_id, f"Unreviewed target exposed: {field.coordinate}"
    assert contract.version == registration.VERSION


@pytest.mark.parametrize("template_id,coordinate,semantic_id", [
    ("emr", "G2", "customer.registration"), ("emr", "G3", "customer.address"),
    ("protocols", "B2", "contractor.registration"), ("protocols", "B3", "contractor.address"),
    ("protocols", "F1", "customer.name"), ("protocols", "F2", "customer.registration"),
    ("protocols", "F3", "customer.address"), ("avk", "F2", "customer.registration"),
    ("avk", "F3", "customer.address"), ("ojr", "F2", "customer.registration"),
    ("ojr", "F3", "customer.address"),
])
def test_only_exact_role_bound_lookup_precedents_are_pdf_writable(
    template_id: str, coordinate: str, semantic_id: str,
) -> None:
    catalog = TemplateCatalog(ROOT, ROOT / "templates/fill-contracts", ROOT / "templates/approved")
    contract = catalog.get(template_id)
    field = contract.field_map[("Данные организации", coordinate)]
    assert field.manual_reason is None
    assert field.semantic_id == semantic_id
    assert field.evidence_rule == "organization_role_pdf"
    for field in contract.fields:
        if field.sheet == "Данные организации" and field.semantic_id is None:
            assert field.manual_reason


def test_lookup_formula_drift_does_not_reassign_another_organization() -> None:
    source = openpyxl.load_workbook(registration.SOURCE_DIR / "4. АВК.xlsx")
    try:
        source["Данные объект"]["B24"] = "='Данные организации'!G2"
        with pytest.raises(ValueError, match="Изменилась ссылка"):
            registration.object_card_field_overrides("avk", source)
    finally:
        source.close()


def test_aosr_representatives_follow_form_roles_not_nearby_old_values() -> None:
    catalog = TemplateCatalog(ROOT, ROOT / "templates/fill-contracts", ROOT / "templates/approved")
    fields = catalog.get("aosr_vl").field_map
    for start, role in (
        (16, "contractor.construction_control"), (19, "contractor.site_representative"),
        (32, "customer.construction_control"), (35, "customer.other_representative"),
    ):
        for offset, part in enumerate(("position", "name", "authority")):
            field = fields[("Данные объект", f"B{start + offset}")]
            assert field.manual_reason is None
            assert field.semantic_id == f"{role}.{part}"
            assert field.evidence_rule == ("authority_document_pdf" if part == "authority" else "signatory_role_pdf")
            assert "электронной подписи" in field.description
    # These groups have no definite representative role in the printed form.
    for coordinate in ("B14", "B15", "B26", "B27", "B28", "B29", "B30", "B31"):
        assert fields[("Данные объект", coordinate)].manual_reason


def test_unreviewed_registry_rows_remain_closed_without_presuming_pdf_absence() -> None:
    catalog = TemplateCatalog(ROOT, ROOT / "templates/fill-contracts", ROOT / "templates/approved")
    for template_id, sheet_prefix in (("emr", "Пр.15 Реестр"), ("avk", " Журнал АВК"), ("ojr", "Раздел")):
        registry = [field for field in catalog.get(template_id).fields if field.sheet.casefold().startswith(sheet_prefix.casefold())]
        assert registry
        for field in registry:
            if template_id == "avk" and field.semantic_id and field.semantic_id.startswith("avk.materials."):
                assert field.allow_project_basis
                continue
            assert field.manual_reason
            if template_id == "avk" and field.cell[0] in "BEFGHIJ" and 36 <= int(field.cell[1:]) <= 86:
                assert "фактической поставки" in field.manual_reason
                continue
            assert "сопоставление" in field.manual_reason
            assert "отсутств" not in field.manual_reason


def test_empty_ojr_inputs_and_avk_installation_are_not_lost_to_discovery() -> None:
    catalog = TemplateCatalog(
        ROOT, ROOT / "templates" / "fill-contracts", ROOT / "templates" / "approved"
    )
    ojr = catalog.get("ojr")
    for cell in ("B5", "B54"):
        assert ojr.field_map[("Данные объект", cell)].manual_reason is None
    avk = catalog.get("avk")
    assert avk.field_map[("Данные объект", "B9")].semantic_id == (
        "project.installation.substation"
    )
    for template_id, city_cell, code_cell in (
        ("avk", "B39", "B40"), ("ojr", "B53", "B54"), ("protocols", "B53", "B54"),
        ("aosr_vl", "B42", "B43"),
    ):
        fields = catalog.get(template_id).field_map
        assert fields[("Данные объект", city_cell)].semantic_id == "project.city"
        assert fields[("Данные объект", city_cell)].manual_reason
        assert ("Данные объект", city_cell) not in {
            (field["sheet"], field["cell"]) for field in catalog.get(template_id).model_fields()
        }
        assert fields[("Данные объект", code_cell)].semantic_id == "project.design_document_code"

    # The source EMR lower block has shifted labels: keep it closed pending a
    # separate layout repair instead of guessing from old project values.
    for cell in ("B54", "B55", "B56"):
        assert catalog.get("emr").field_map[("Данные объект", cell)].manual_reason


def test_project_basis_is_restricted_to_reviewed_material_and_work_quantities() -> None:
    catalog = TemplateCatalog(ROOT, ROOT / "templates/fill-contracts", ROOT / "templates/approved")
    expected_counts = {"aosr_vl": 11, "emr": 69, "avk": 102, "protocols": 0, "ojr": 0}
    for template_id, expected in expected_counts.items():
        contract = catalog.get(template_id)
        project_fields = [field for field in contract.fields if field.allow_project_basis]
        assert len(project_fields) == expected
        assert contract.structural_findings["project_basis_target_count"] == expected
        for field in project_fields:
            assert field.manual_reason is None
            assert field.evidence_rule == "actual_executive_document_only"
            assert field.value_kind != "date"
            assert field.semantic_id
            assert not any(word in field.semantic_id for word in ("quality_document", "authority", "number", "test_conditions"))
            assert "value_basis=project" in field.description
        assert contract.approved is False
        assert contract.status == "DISCOVERY_REVIEW_REQUIRED"
    aosr = catalog.get("aosr_vl")
    assert {(field.sheet, field.cell) for field in aosr.fields if field.allow_project_basis} == (
        registration.AOSR_VL_PROJECT_QUANTITY_TARGETS
    )


def test_emr_material_rows_preserve_item_identity_and_do_not_open_serial_numbers() -> None:
    catalog = TemplateCatalog(ROOT, ROOT / "templates/fill-contracts", ROOT / "templates/approved")
    fields = catalog.get("emr").field_map
    for row in range(19, 42):
        for column, part in (("E", "name"), ("AI", "type"), ("BE", "quantity")):
            field = fields[("Ведомость общ", f"{column}{row}")]
            assert field.semantic_id == f"emr.materials.item_{row - 18}.{part}"
            assert field.allow_project_basis
            assert field.required is False
            assert "одной позиции" in field.description
        if ("Ведомость общ", f"AU{row}") in fields:
            assert fields[("Ведомость общ", f"AU{row}")].manual_reason


def test_avk_uses_all_51_independent_rows_without_old_act_provenance() -> None:
    catalog = TemplateCatalog(ROOT, ROOT / "templates/fill-contracts", ROOT / "templates/approved")
    fields = catalog.get("avk").field_map
    for row in range(36, 87):
        for column, part in (("C", "name"), ("D", "quantity")):
            field = fields[(" Журнал АВК", f"{column}{row}")]
            assert field.semantic_id == f"avk.materials.item_{row - 35}.{part}"
            assert field.allow_project_basis
        for column in "BEFGHIJ":
            field = fields[(" Журнал АВК", f"{column}{row}")]
            assert field.manual_reason
            assert not field.allow_project_basis
    for coordinate in ("C43", "G43", "A41"):
        assert fields[("стойки", coordinate)].manual_reason


def test_avk_journal_preserves_headers_geometry_and_only_empty_inputs() -> None:
    catalog = TemplateCatalog(ROOT, ROOT / "templates/fill-contracts", ROOT / "templates/approved")
    contract = catalog.get("avk")
    source = openpyxl.load_workbook(registration.SOURCE_DIR / "4. АВК.xlsx")
    candidate = openpyxl.load_workbook(catalog.candidate_path(contract))
    try:
        assert candidate.sheetnames == source.sheetnames
        before, after = source[" Журнал АВК"], candidate[" Журнал АВК"]
        assert str(after.print_area) == str(before.print_area)
        assert str(after.merged_cells) == str(before.merged_cells)
        assert after.page_setup == before.page_setup
        assert after.page_margins == before.page_margins
        for row in (34, 35):
            for column in "ABCDEFGHIJ":
                coordinate = f"{column}{row}"
                assert after[coordinate].value == before[coordinate].value
                assert (after.title, coordinate) not in contract.field_map
        for row in range(36, 87):
            assert after[f"A{row}"].value == row - 35
            for column in "BCDEFGHIJ":
                coordinate = f"{column}{row}"
                assert after[coordinate].value is None
                assert (after.title, coordinate) in contract.field_map
                assert after[coordinate]._style == before[coordinate]._style
        assert {f.sheet for f in contract.fields if f.allow_project_basis} == {" Журнал АВК"}
    finally:
        candidate.close()
        source.close()


def test_avk_actual_columns_reject_header_and_row_structure_drift() -> None:
    source = openpyxl.load_workbook(registration.SOURCE_DIR / "4. АВК.xlsx")
    try:
        assert len(registration.avk_journal_layout_changes(source)) == 51 * 7
        source[" Журнал АВК"]["E34"] = "Проектная организация"
        with pytest.raises(ValueError, match="Изменилась подпись"):
            registration.avk_journal_layout_changes(source)
        source[" Журнал АВК"]["E34"] = "Поставщик"
        source[" Журнал АВК"]["A36"] = 99
        with pytest.raises(ValueError, match="Изменилась строка"):
            registration.avk_journal_layout_changes(source)
    finally:
        source.close()


def test_avk_project_row_marks_only_materials_and_keeps_actual_columns_blank(tmp_path: Path) -> None:
    from executive_docs.domain import TemplateCellAssignment
    from executive_docs.selected_templates import PROJECT_FILL_RGB, SelectedTemplateGenerator

    catalog = TemplateCatalog(ROOT, ROOT / "templates/fill-contracts", ROOT / "templates/approved")
    contract = catalog.get("avk")
    fragment = "Проект: тестовый материал, 12 шт."
    assignments = [
        TemplateCellAssignment(
            sheet=" Журнал АВК", cell=cell, value=value,
            source_file_id="offline-test-pdf", locator="page:1",
            evidence_fragment=fragment, value_basis="project",
        )
        for cell, value in (("C36", "тестовый материал"), ("D36", "12 шт."))
    ]
    output, unresolved = SelectedTemplateGenerator(catalog).generate(contract, assignments, tmp_path)
    book = openpyxl.load_workbook(output)
    try:
        journal = book[" Журнал АВК"]
        for column in "CD":
            assert journal[f"{column}36"].fill.fgColor.rgb == PROJECT_FILL_RGB
        for row in range(36, 87):
            for column in "BEFGHIJ":
                cell = journal[f"{column}{row}"]
                assert cell.value is None
                assert cell.fill.fgColor.rgb == contract.warning_fill_rgb
        assert "ЧЕРНОВИК" in journal.oddHeader.left.text
        assert "Факт выполнения не подтверждён" in journal.oddHeader.left.text
        assert all(item.category == "manual_confirmation" for item in unresolved if item.sheet == journal.title and item.cell[0] in "BEFGHIJ" and 36 <= int(item.cell[1:]) <= 86)
    finally:
        book.close()


@pytest.mark.parametrize("template_id,filename,coordinate", [
    ("emr", "1. ЭМР1.xlsx", "AI17"), ("avk", "4. АВК.xlsx", "D34"),
])
def test_material_column_mapping_rejects_source_header_drift(
    template_id: str, filename: str, coordinate: str,
) -> None:
    source = openpyxl.load_workbook(registration.SOURCE_DIR / filename)
    try:
        assert registration.material_table_field_overrides(template_id, source)
        sheet = "Ведомость общ" if template_id == "emr" else " Журнал АВК"
        source[sheet][coordinate] = "Фактический результат измерения"
        with pytest.raises(ValueError, match="Изменилась подпись"):
            registration.material_table_field_overrides(template_id, source)
    finally:
        source.close()


def test_protocol_rows_with_stale_measurements_are_not_newly_opened() -> None:
    catalog = TemplateCatalog(ROOT, ROOT / "templates/fill-contracts", ROOT / "templates/approved")
    fields = catalog.get("protocols").field_map
    for sheet, coordinate in (("прот.№6", "C13"), ("прот.№7", "B23"), ("прот.№8", "D15")):
        assert fields[(sheet, coordinate)].manual_reason
        assert fields[(sheet, coordinate)].allow_project_basis is False
